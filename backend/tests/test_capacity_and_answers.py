"""
Round-one fixes from the end-to-end run: a turn fits a free-tier model's
minute, an answer that outgrows a small model moves to a bigger one, what the
reader sees is the answer and nothing else, and retrieval reaches the passage
that holds the fact. None of these call a model.
"""
import asyncio
import json
from types import SimpleNamespace

import pytest


# ── What people read when capacity runs out ──────────────────────────────────

@pytest.mark.parametrize("raw", [
    "Agent has reached an unrecoverable state due to max_tokens limit.",
    "Error code: 429 - Rate limit reached for model `qwen/qwen3.6-27b` on input tokens per minute (ITPM): Limit 7000",
    "Error code: 429 - Rate limit reached for model `openai/gpt-oss-120b` on tokens per day (TPD): Limit 200000",
])
def test_capacity_failures_read_without_limit_wording_or_settings(raw):
    from core.errors import humanise
    msg = humanise(raw)
    assert "limit" not in msg.lower() and "MAX_OUTPUT_TOKENS" not in msg and "LLM_BACKEND" not in msg
    assert "qwen" not in msg and "gpt-oss" not in msg


def test_a_refresh_that_only_lacked_model_time_is_not_shown_as_an_error():
    from core.errors import refresh_error
    assert refresh_error("Error code: 429 - Rate limit reached ... tokens per minute") is None
    assert refresh_error("Agent has reached an unrecoverable state due to max_tokens limit.") is None
    assert refresh_error("tokens per day (TPD)") is None
    assert "brand new" in refresh_error("Some brand new failure")


# ── What the reader sees is the answer ───────────────────────────────────────

def test_leaked_reasoning_is_removed_from_answers():
    from core.formatting import plain_answer
    leaked = ("KAN-1 is titled sprint plan.\n\nI can pull the live version.\n</think>\n\n"
              "KAN-1 is titled **sprint plan**; status **Idea**, priority **Medium**.")
    assert plain_answer(leaked) == "KAN-1 is titled **sprint plan**; status **Idea**, priority **Medium**."
    assert plain_answer("<think>hmm, the runbook</think>Roll back in Argo CD.") == "Roll back in Argo CD."
    assert plain_answer("Answer first.<think>unfinished reasoning") == "Answer first."


def test_lookalike_hyphens_and_spaces_become_plain_ones():
    from core.formatting import plain_answer
    out = plain_answer("Change CHG‑5518 lifts uptime to 99.0 % (was 50 %).")
    assert out == "Change CHG-5518 lifts uptime to 99.0 % (was 50 %)."


def _run_gate(script):
    from chat.routes import AnswerGate
    gate, events = AnswerGate(), []
    for kind, value in script:
        events += gate.feed(value) if kind == "text" else gate.discard() if kind == "tool" else []
    events += gate.finish()
    return gate, events


def test_narration_before_a_tool_call_never_reaches_the_reader():
    gate, events = _run_gate([("text", "I need to fetch the open Jira tickets first."), ("tool", None),
                              ("text", "There is one open ticket, **KAN-1**.")])
    assert [e["type"] for e in events] == ["text"]
    assert gate.answer == "There is one open ticket, **KAN-1**."


def test_text_already_shown_is_withdrawn_when_a_tool_call_follows():
    long_narration = "KAN-1 is the sprint plan. " * 20
    gate, events = _run_gate([("text", long_narration), ("tool", None), ("text", "Live: KAN-1 is Idea.")])
    kinds = [e["type"] for e in events]
    assert kinds[0] == "text" and "reset" in kinds and kinds[-1] == "text"
    assert gate.answer == "Live: KAN-1 is Idea."


def test_reasoning_closed_mid_stream_is_taken_back():
    reasoning = "Let me think about which passage covers the rollback steps. " * 10
    gate, events = _run_gate([("text", reasoning), ("text", "</think>Argo CD → History → Rollback to previous.")])
    shown = ""
    for e in events:
        shown = "" if e["type"] == "reset" else shown + e["delta"]
    assert shown == gate.answer.lstrip() == "Argo CD → History → Rollback to previous."


# ── A turn fits in a model's minute ──────────────────────────────────────────

def test_qwen_requests_switch_reasoning_off_and_keep_their_small_budget():
    from agent.agent import request_params
    qwen = request_params("qwen/qwen3.6-27b", background=False)
    assert qwen["max_tokens"] == 900 and qwen["extra_body"]["reasoning_effort"] == "none"
    # Top-level unknown keywords are refused by the OpenAI client itself.
    assert "reasoning_format" not in qwen and "reasoning_effort" not in qwen
    assert "extra_body" not in request_params("openai/gpt-oss-120b", background=False)


def test_background_reports_do_not_use_models_that_cannot_write_them(monkeypatch):
    from agent import agent as a
    for key in ("CEREBRAS_API_KEY", "GEMINI_API_KEY", "OPENROUTER_API_KEY"):
        monkeypatch.setattr(a.settings, key, "")
    monkeypatch.setattr(a.settings, "GROQ_API_KEY", "g")
    monkeypatch.setattr(a.settings, "GROQ_MODEL", "openai/gpt-oss-120b")
    monkeypatch.setattr(a.settings, "GROQ_BACKGROUND_MODEL", "openai/gpt-oss-20b")
    monkeypatch.setattr(a.settings, "GROQ_FALLBACK_MODELS", "openai/gpt-oss-120b,qwen/qwen3.8-27b,qwen/qwen3.6-27b")
    assert [c[3] for c in a.model_chain(background=True)] == ["openai/gpt-oss-20b", "openai/gpt-oss-120b"]
    assert "qwen/qwen3.6-27b" in [c[3] for c in a.model_chain(background=False)]


def test_tools_are_offered_only_to_questions_that_need_them():
    from agent.agent import wants_file, wants_jira, wants_write
    assert wants_file("Write a one-page handover note as a Word document")
    assert wants_file("Put the open KAN tickets in a spreadsheet")
    assert not wants_file("How do I roll back dispatch-api in production?")
    assert wants_jira("What is Jira issue KAN-1 about?") and wants_jira("Which tickets are blocked?")
    assert not wants_jira("How do I run apollo-delivery-service locally?")
    assert not wants_jira("What was fixed in dispatch-api release 2.3.1?")
    assert wants_write("Create an issue for the quota problem")
    assert not wants_write("List the open issues and pull requests")


def test_a_huge_json_tool_result_is_compacted_to_what_the_model_needs():
    from agent.agent import compact_tool_text
    issues = [{"number": i, "title": f"Issue {i}", "state": "open", "node_id": "I_kw" * 10,
               "html_url": f"https://github.com/o/r/issues/{i}", "url": "https://api.github.com/x" * 3,
               "user": {"login": "priya", "avatar_url": "https://avatars" * 5, "id": 7, "type": "User"},
               "reactions": {"total_count": 0, "+1": 0}, "body": "details " * 30}
              for i in range(40)]
    out = compact_tool_text(json.dumps(issues), limit=6000)
    assert len(out) <= 6100
    assert "Issue 0" in out and "priya" in out
    assert "avatar" not in out and "node_id" not in out and "reactions" not in out


def test_the_trimmer_rewrites_only_long_results():
    from agent.agent import ToolResultTrimmer
    trimmer = ToolResultTrimmer(limit=100)
    short = SimpleNamespace(result={"toolUseId": "1", "status": "success", "content": [{"text": "fine"}]})
    trimmer._after_tool(short)
    assert short.result["content"] == [{"text": "fine"}]
    long = SimpleNamespace(result={"toolUseId": "2", "status": "success", "content": [{"text": "x\n" * 500}]})
    trimmer._after_tool(long)
    assert len(long.result["content"][0]["text"]) < 200


# ── An answer that outgrows a small model moves to a bigger one ──────────────

def _outgrown(a, monkeypatch, chain, current, exhausted=()):
    a._EXHAUSTED.clear()
    for key in exhausted:
        a.mark_exhausted(key, 3600)
    monkeypatch.setattr(a.settings, "LLM_BACKEND", "groq")
    monkeypatch.setattr(a, "model_chain", lambda background: chain)
    monkeypatch.setattr(a, "_build_model", lambda background=False, entry=None: SimpleNamespace(sensei_key=entry[0]))
    hook = a.ModelFailover(max_wait_s=0.2)
    notices = []
    hook.on_notice = notices.append
    event = SimpleNamespace(retry=False, exception=None,
                            stop_response=SimpleNamespace(stop_reason="max_tokens"),
                            agent=SimpleNamespace(model=SimpleNamespace(sensei_key=current)))
    asyncio.run(hook._after_model_call(event))
    a._EXHAUSTED.clear()
    return event, notices


CHAIN = [("groq/openai/gpt-oss-120b", "groq", "u", "openai/gpt-oss-120b"),
         ("groq/qwen/qwen3.6-27b", "groq", "u", "qwen/qwen3.6-27b"),
         ("groq/openai/gpt-oss-20b", "groq", "u", "openai/gpt-oss-20b")]


def test_an_answer_too_long_for_qwen_is_redone_on_a_model_with_room(monkeypatch):
    from agent import agent as a
    event, notices = _outgrown(a, monkeypatch, CHAIN, "groq/qwen/qwen3.6-27b", exhausted=["groq/openai/gpt-oss-120b"])
    assert event.retry is True
    assert event.agent.model.sensei_key == "groq/openai/gpt-oss-20b"
    assert notices == ["Taking a little longer to check"]


def test_an_answer_too_long_for_a_roomy_model_is_not_bounced_around(monkeypatch):
    from agent import agent as a
    event, _ = _outgrown(a, monkeypatch, CHAIN, "groq/openai/gpt-oss-120b")
    assert event.retry is False


def test_a_provider_that_rejects_the_reasoning_switches_is_retried_without_them():
    from agent import agent as a

    class Model:
        sensei_key = "groq/qwen/qwen3.6-27b"

        def __init__(self):
            self.config = {"params": {"max_tokens": 900,
                                      "extra_body": {"reasoning_effort": "none", "reasoning_format": "hidden"}}}

        def update_config(self, **kw):
            self.config.update(kw)

    model = Model()
    event = SimpleNamespace(retry=False, stop_response=None, agent=SimpleNamespace(model=model),
                            exception=Exception("Error code: 400 - 'reasoning_effort' is not supported with this model"))
    asyncio.run(a.ModelFailover()._after_model_call(event))
    assert event.retry is True and model.config["params"] == {"max_tokens": 900}


def test_a_client_that_refuses_the_reasoning_keywords_is_retried_without_them():
    from agent import agent as a
    assert a._rejects_reasoning_params(
        TypeError("AsyncCompletions.create() got an unexpected keyword argument 'reasoning_format'"))
    assert not a._rejects_reasoning_params(TypeError("unexpected keyword argument 'foo'"))


# ── Retrieval reaches the passage that holds the fact ────────────────────────

def _row(cid, title, distance, label="Strands README"):
    return (cid, "doc", {"source_label": label, "title": title}, distance)


def test_one_long_page_still_fills_every_place():
    from agent.tools import interleave
    only = [_row(f"r{i}", "README", 0.1 * i) for i in range(8)]
    assert [r[0] for r in interleave([only, only], 5, per_page=2)] == ["r0", "r1", "r2", "r3", "r4"]


class _Collection:
    """A tiny stand-in for a Chroma collection that honours $contains."""

    def __init__(self, docs):
        self.docs = docs     # id -> text

    def count(self):
        return len(self.docs)

    def get(self, where_document=None, include=None, limit=None, where=None):
        term = (where_document or {}).get("$contains", "")
        ids = [i for i, d in self.docs.items() if term in d]
        return {"ids": ids[:limit] if limit else ids}


def test_distinctive_terms_prefer_versions_and_rare_words():
    from agent.tools import distinctive_terms
    docs = {f"c{i}": "dispatch-api release notes and deployment" for i in range(60)}
    docs["changelog"] = "## 2.3.1 — Fix: webhook retry used the old secret"
    docs["readme"] = "## Running locally\npip install -r requirements.txt"
    col = _Collection(docs)
    assert distinctive_terms("What was fixed in dispatch-api release 2.3.1?", col, None)[0] == "2.3.1"
    assert "locally" in distinctive_terms("How do I run apollo-delivery-service locally?", col, None)
    assert "release" not in distinctive_terms("What was fixed in dispatch-api release 2.3.1?", col, None)


def test_citations_are_per_page_not_per_space(monkeypatch):
    import db.chroma
    from agent.tools import make_search_tool

    class Collection:
        def count(self):
            return 2

        def query(self, query_texts, n_results, include, where=None, where_document=None):
            rows = [("a", "[Confluence: SD › Decision log]\nWe chose Postgres", {"source_label": "Confluence: SD", "title": "Decision log"}, 0.2),
                    ("b", "[Confluence: SD › On-call]\nPriya covers incidents", {"source_label": "Confluence: SD", "title": "On-call"}, 0.3)]
            return {"ids": [[r[0] for r in rows]], "documents": [[r[1] for r in rows]],
                    "metadatas": [[r[2] for r in rows]], "distances": [[r[3] for r in rows]]}

        def get(self, **kw):
            return {"ids": []}

    monkeypatch.setattr(db.chroma, "get_workspace_collection", lambda client, ws: Collection())
    search, captured = make_search_tool("ws", None, n_results=2)
    search.unmetered("who covers incidents")
    assert [c["title"] for c in captured] == ["Decision log", "On-call"]
    assert captured[1]["excerpt"].startswith("Priya covers incidents")


# ── Files and pages are read whole ───────────────────────────────────────────

def test_word_tables_headers_and_footers_are_read(tmp_path):
    from docx import Document
    from agent.tools import docx_text
    doc = Document()
    doc.add_paragraph("SwiftMaps renewal")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Term", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "Annual fee", "GBP 48,000"
    doc.sections[0].header.paragraphs[0].text = "CONFIDENTIAL"
    doc.sections[0].footer.paragraphs[0].text = "Owner: Priya Nair"
    path = tmp_path / "renewal.docx"
    doc.save(path)
    text = docx_text(str(path))
    assert "Annual fee | GBP 48,000" in text
    assert "CONFIDENTIAL" in text and "Owner: Priya Nair" in text
    assert text.index("SwiftMaps renewal") < text.index("Annual fee")


def test_a_plain_text_page_is_named_by_its_heading_or_file():
    from bs4 import BeautifulSoup
    from agent.tools import page_title
    md = "# Strands Agents\n\nBoth SDKs default to Amazon Bedrock."
    assert page_title(BeautifulSoup(md, "html.parser"), md, "https://raw.example.com/sdk/README.md") == "Strands Agents"
    assert page_title(BeautifulSoup("plain", "html.parser"), "plain", "https://raw.example.com/sdk/README.md") == "README.md"


# ── Page data and meetings ───────────────────────────────────────────────────

def test_readiness_sources_lose_their_citation_markers():
    from agent.readiness import clean_source
    assert clean_source("[5] Confluence: SD – Apollo — Deployment runbook") == "Confluence: SD – Apollo — Deployment runbook"
    assert clean_source("[2] .github/workflows/ci.yml") == ".github/workflows/ci.yml"


def test_sensei_answers_when_addressed_even_if_it_could_not_think():
    from meetings.listener import fallback_reply
    addressed = fallback_reply("Sensei, who owns the routing worker?", "busy")
    assert addressed.kind == "answer" and "again" in addressed.text and "busy" not in addressed.text
    assert fallback_reply("We deploy on Fridays", "busy").kind == "silent"


def test_the_summary_sees_what_sensei_answered():
    from meetings.listener import transcript_with_replies
    text = transcript_with_replies(
        [{"speaker": "Ravi", "text": "Sensei, who owns the routing worker?"}],
        [{"kind": "answer", "text": "Daniel Okafor.", "trigger": "Sensei, who owns the routing worker?", "confidence": 1.0}])
    assert "Sensei (answering, from the sources): Daniel Okafor." in text


class _WordCollection(_Collection):
    """Honours $contains and $or of them, the way the ranking asks."""

    def get(self, where_document=None, include=None, limit=None, where=None):
        clauses = (where_document or {}).get("$or") or [where_document or {}]
        terms = [c.get("$contains", "") for c in clauses]
        ids = [i for i, d in self.docs.items() if any(t in d for t in terms)]
        return {"ids": ids[:limit] if limit else ids}


def test_a_chunk_holding_the_questions_rare_words_outranks_a_vaguely_similar_one():
    from agent.tools import lexical_match, term_weights
    docs = {f"n{i}": "Strands agents docs: model providers and tools overview" for i in range(30)}
    docs["readme"] = "Both SDKs default to the Amazon Bedrock model provider.\npip install strands-agents strands-agents-tools"
    col = _WordCollection(docs)
    weights = term_weights("Which model provider is the default, and what pip command installs the tools package?", col, None)
    assert "default" in weights and "model" not in weights       # a word in every chunk tells nothing
    assert lexical_match(docs["readme"], weights) > lexical_match(docs["n0"], weights)
    assert term_weights("anything", SimpleNamespace(), None) == {}   # a store that cannot count ranks by meaning alone


def test_a_page_keeps_a_third_place_when_every_other_page_is_far_behind():
    from agent.tools import pick_diverse

    def row(cid, title):
        return (cid, "doc", {"source_label": "Strands SDK docs", "title": title}, 0.0)

    readme = [(0.79, row("r2", "README")), (0.73, row("r7", "README")), (0.70, row("r3", "README")),
              (0.69, row("r5", "README"))]
    other = [(0.52, row("a38", "Architecture")), (0.66, row("l10", "Agent Loop"))]
    picked = [r[0] for r in pick_diverse(readme + other, 5, per_page=2, margin=0.08)]
    assert picked == ["r2", "r7", "l10", "r3", "r5"]


def test_stems_match_the_forms_sources_use():
    from agent.tools import _stem
    assert _stem("fixed") == "fix" and _stem("installs") == "install" and _stem("locally") == "locally"


def test_only_the_pages_an_answer_drew_on_are_shown_as_its_sources():
    from agent.tools import relevant_citations
    cites = [
        {"index": 1, "source_label": "Confluence: SD", "title": "Apollo — Deployment runbook",
         "excerpt": "Rollback", "_text": "Argo CD → application dispatch-api → History → Rollback to previous. Takes about 90 seconds."},
        {"index": 2, "source_label": "leeds_depot_cutover_plan.txt", "title": "leeds_depot_cutover_plan.txt",
         "excerpt": "Cut-over", "_text": "Cut-over night is Sat 14 Nov 2026; gate scanners arrive 6 November."},
    ]
    answer = "Use Argo CD: open dispatch-api, go to History and choose Rollback to previous. It takes about 90 seconds."
    kept = relevant_citations(cites, answer)
    assert [c["title"] for c in kept] == ["Apollo — Deployment runbook"]
    assert all(not k.startswith("_") for c in kept for k in c)
    # An answer sharing nothing with its passages still shows where it searched.
    assert len(relevant_citations(cites, "I could not find that.")) == 1
    assert relevant_citations([], "anything") == []


def test_one_shared_word_does_not_make_a_granted_tool_relevant():
    from toolgrants.registry import select_relevant
    tools = [SimpleNamespace(tool_name="github_list_commits", mcp_tool=SimpleNamespace(description="List commits in a Python repository")),
             SimpleNamespace(tool_name="github_list_issues", mcp_tool=SimpleNamespace(description="List issues in a GitHub repository")),
             SimpleNamespace(tool_name="github_list_pull_requests", mcp_tool=SimpleNamespace(description="List pull requests"))] + \
            [SimpleNamespace(tool_name=f"github_other_{i}", mcp_tool=SimpleNamespace(description="Unrelated")) for i in range(6)]
    readme = {t.tool_name for t in select_relevant(tools, "Which pip command installs the Python SDK?", limit=4)}
    assert readme == set()
    listing = {t.tool_name for t in select_relevant(tools, "List the open issues and pull requests in apollo-delivery-service.", limit=4)}
    assert {"github_list_issues", "github_list_pull_requests"} <= listing


def test_page_data_saved_with_old_wording_is_served_in_todays():
    from db.models import serialize_brief, serialize_gap_report
    brief = serialize_brief({"_id": "b", "workspace_id": "w", "user_id": "u", "status": "error",
                             "error_message": "The answer grew past the model's output limit. Ask for something narrower, or raise MAX_OUTPUT_TOKENS."})
    assert "MAX_OUTPUT_TOKENS" not in brief["error_message"] and "limit" not in brief["error_message"].lower()
    kept = serialize_brief({"_id": "b", "workspace_id": "w", "user_id": "u", "status": "ready", "brief": {},
                            "refresh_error": "The model provider's daily limit was reached during the last refresh."})
    assert kept["refresh_error"] is None
    report = serialize_gap_report({"_id": "g", "workspace_id": "w", "status": "ready", "gaps": [],
                                   "refresh_error": "Could not reach a service we depend on."})
    assert report["refresh_error"] == "Could not reach a service we depend on."   # a real failure stays visible
